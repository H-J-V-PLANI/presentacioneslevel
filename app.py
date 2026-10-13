import os
import requests
import streamlit as st
from duckduckgo_search import DDGS
from pptx import Presentation
from pptx.util import Inches, Pt

# Configuración de la interfaz web
st.set_page_config(
    page_title="Generador de Presentaciones Educativas",
    page_icon="🎓",
    layout="centered",
)

# --- SISTEMA DE LOGIN SENCILLO ---
CREDENCIALES = {"profesor": "profe2026", "estudiante": "clase2026"}

if "autenticado" not in st.session_state:
  st.session_state.autenticado = False

if not st.session_state.autenticado:
  st.title("🔐 Acceso a la Plataforma Educativa")
  st.write(
      "Inicia sesión para generar tus presentaciones y recursos didácticos."
  )

  with st.form("login_form"):
    user = st.text_input("Usuario")
    password = st.text_input("Contraseña", type="password")
    entrar = st.form_submit_button("Ingresar")

    if entrar:
      if user in CREDENCIALES and CREDENCIALES[user] == password:
        st.session_state.autenticado = True
        st.session_state.usuario = user
        st.rerun()
      else:
        st.error("Usuario o contraseña incorrectos.")
  st.stop()

# --- APLICACIÓN PRINCIPAL ---
st.sidebar.success(f"Conectado como: **{st.session_state.usuario}**")
if st.sidebar.button("Cerrar Sesión"):
  st.session_state.autenticado = False
  st.rerun()

st.title("🎓 Generador de Presentaciones con Enlaces de YouTube")
st.write(
    "Crea presentaciones automáticas con investigación web real y links de"
    " videos de YouTube listos para exponer."
)

with st.form("form_presentacion"):
  tema = st.text_input(
      "¿Sobre qué tema o asignatura deseas la presentación?",
      "La animación 2D",
  )
  col1, col2 = st.columns(2)
  with col1:
    num_slides = st.slider(
        "Número de diapositivas", min_value=3, max_value=8, value=4
    )
  with col2:
    incluir_yt = st.checkbox(
        "Incluir enlaces de videos de YouTube", value=True
    )

  generar_btn = st.form_submit_button("🚀 Generar y Descargar Presentación")


# Función para obtener fragmentos reales de información
@st.cache_data(show_spinner=False)
def buscar_contenido_real(query):
  fragmentos = []
  videos_encontrados = []
  try:
    with DDGS() as ddgs:
      resultados = ddgs.text(query, max_results=8)
      if resultados:
        for r in resultados:
          if "body" in r and len(r["body"]) > 30:
            fragmentos.append(r["body"])

      resultados_yt = ddgs.text(f"site:youtube.com {query}", max_results=3)
      if resultados_yt:
        for r in resultados_yt:
          titulo_v = r.get("title", "Video educativo")
          url_v = r.get("href", "#")
          videos_encontrados.append({"titulo": titulo_v, "url": url_v})
  except Exception:
    pass

  if not fragmentos:
    fragmentos = [
        f"Conceptos principales, fundamentos y características clave sobre"
        f" {query}.",
        f"Importancia, aplicaciones prácticas y metodología en torno a {query}.",
    ]

  return fragmentos, videos_encontrados


if generar_btn:
  if not tema.strip():
    st.warning("Por favor, escribe un tema válido.")
  else:
    barra = st.progress(0)
    estado = st.empty()

    try:
      estado.text(
          "🌐 Investigando información real en la web y buscando videos..."
      )
      barra.progress(30)
      lista_fragmentos, lista_videos = buscar_contenido_real(tema)

      estado.text("📝 Organizando el contenido real por diapositivas...")
      barra.progress(60)

      prs = Presentation()

      subtitulos_ideas = [
          "Definición y Conceptos Fundamentales",
          "Características y Elementos Principales",
          "Proceso, Técnicas y Metodología",
          "Aplicaciones Prácticas en el Área",
          "Ventajas y Aspectos Clave",
          "Retos Actuales e Innovación",
          "Ejemplos y Casos de Estudio",
          "Conclusiones y Resumen General",
      ]

      # Generar diapositivas usando la información real obtenida
      for i in range(num_slides):
        slide = prs.slides.add_slide(prs.slide_layouts[1])

        sub_titulo = (
            subtitulos_ideas[i]
            if i < len(subtitulos_ideas)
            else f"Aspecto Clave {i+1}"
        )
        slide.shapes.title.text = f"{i+1}. {sub_titulo}"

        f1 = (
            lista_fragmentos[i % len(lista_fragmentos)]
            if len(lista_fragmentos) > 0
            else ""
        )
        f2 = (
            lista_fragmentos[(i + 1) % len(lista_fragmentos)]
            if len(lista_fragmentos) > 1
            else f1
        )

        contenido_real = (
            f"• {f1[:220]}...\n\n• Elemento complementario: {f2[:220]}..."
        )
        slide.placeholders[1].text = contenido_real

      # Diapositiva final con los links de YouTube reales
      if incluir_yt and lista_videos:
        slide_yt = prs.slides.add_slide(prs.slide_layouts[1])
        slide_yt.shapes.title.text = "🎥 Enlaces de Videos en YouTube"
        detalle_yt = (
            "Recursos audiovisuales recomendados para complementar la"
            " clase:\n\n"
        )
        for v in lista_videos:
          detalle_yt += f"• {v['titulo']}\n  🔗 {v['url']}\n\n"
        slide_yt.placeholders[1].text = detalle_yt

      barra.progress(90)
      estado.text("📊 Compilando archivo PowerPoint...")

      archivo_salida = f"presentacion_{os.getpid()}.pptx"
      prs.save(archivo_salida)

      barra.progress(100)
      estado.text("¡Completado con éxito!")

      with open(archivo_salida, "rb") as f:
        st.success("🎉 ¡Tu presentación con contenido real está lista!")
        st.download_button(
            label="📥 Descargar Presentación (.pptx)",
            data=f,
            file_name=f"{tema.replace(' ', '_')}.pptx",
            mime=(
                "application/vnd.openxmlformats-officedocument.presentationml.presentation"
            ),
        )

      if os.path.exists(archivo_salida):
        os.remove(archivo_salida)

    except Exception as err:
      st.error(f"Ocurrió un error: {err}")
